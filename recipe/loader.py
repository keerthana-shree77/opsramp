"""Read a recipe file (CSV / XLSX / JSON) into raw, un-interpreted rows.

The loader only deals with *shape*: locating the header, mapping column
spellings onto canonical field names and preserving the original cell values.
Meaning is added later by :mod:`recipe.validator` and :mod:`recipe.normalizer`.
Uploaded files are parsed as data - never imported, evaluated or executed.
"""
from __future__ import annotations

import csv
import io
import json
import logging
import re
from dataclasses import dataclass, field
from pathlib import Path

from firmware.normalizer import norm_text

logger = logging.getLogger(__name__)

CANONICAL_FIELDS = (
    "platform",
    "model",
    "category",
    "component",
    "target_version",
    "target_build",
    "mandatory",
    "applies_to_family",
    "notes",
)

REQUIRED_FIELDS = ("platform", "category", "component", "target_version")

COLUMN_ALIASES: dict[str, str] = {}


def _register(field_name: str, *spellings: str) -> None:
    for spelling in spellings:
        COLUMN_ALIASES[norm_text(spelling)] = field_name


_register("platform", "platform", "vendor platform", "platform family", "family",
          "vendor", "make", "manufacturer", "product family", "hw platform",
          "hardware platform", "server platform", "system family")
_register("model", "model", "model name", "hardware model", "product", "product name",
          "device model", "type", "server model", "hw model", "machine type",
          "system model")
_register("category", "category", "infrastructure category", "asset category",
          "device category", "class", "device type", "asset type", "hardware type",
          "infra category", "infrastructure type")
_register("component", "component", "firmware component", "component name", "part",
          "firmware", "software component", "item", "fw component", "firmware name",
          "firmware type", "component type")
_register("target_version", "target version", "target", "latest version",
          "approved version", "recipe version", "required version", "version",
          "expected version", "baseline version", "target firmware version",
          "recommended version", "target fw version", "latest fw version",
          "fw version", "firmware version", "baseline", "golden version",
          "desired version", "to version")
_register("target_build", "target build", "build", "build number", "target build number",
          "approved build")
_register("mandatory", "mandatory", "required", "is mandatory", "compliance mandatory",
          "enforced")
_register("applies_to_family", "applies to family", "family match", "model family",
          "apply to family", "wildcard")
_register("notes", "notes", "comment", "comments", "remarks", "description")

MAX_ROWS = 20000
MAX_HEADER_SCAN = 25


@dataclass
class RawRecipe:
    rows: list[dict] = field(default_factory=list)
    columns: list[str] = field(default_factory=list)
    mapped_columns: dict[str, str] = field(default_factory=dict)
    source_name: str = ""
    sheet: str = ""
    layout: str = "table"
    notes: list[str] = field(default_factory=list)

    @property
    def missing_required(self) -> list[str]:
        present = set(self.mapped_columns.values())
        return [f for f in REQUIRED_FIELDS if f not in present]


class RecipeLoadError(Exception):
    """A recipe file could not be read at all."""


def _map_header(header: list[str]) -> dict[str, str]:
    mapping: dict[str, str] = {}
    used: set[str] = set()
    for raw in header:
        key = norm_text(raw)
        if not key:
            continue
        target = COLUMN_ALIASES.get(key)
        if target is None:
            # Tolerate "Target Version (Y-axis)" style decorations.
            for alias, candidate in COLUMN_ALIASES.items():
                if key.startswith(alias) and len(alias) >= 5:
                    target = candidate
                    break
        if target and target not in used:
            mapping[raw] = target
            used.add(target)
    return mapping


def _header_score(row: list[str]) -> int:
    """How many of the required columns this row appears to name."""
    mapping = _map_header([str(c) for c in row])
    return len(set(mapping.values()) & set(REQUIRED_FIELDS))


def _looks_like_header(row: list[str]) -> bool:
    return _header_score(row) >= 2


def find_header(table: list[list]) -> int | None:
    """Index of the best header row within the scan window, or None."""
    best_index = None
    best_score = 1  # a single matching column is too weak to act on
    for index, row in enumerate(table[:MAX_HEADER_SCAN]):
        cells = ["" if c is None else str(c) for c in row]
        if not any(cells):
            continue
        score = _header_score(cells)
        if score > best_score:
            best_score = score
            best_index = index
            if score == len(REQUIRED_FIELDS):
                break
    return best_index


def describe_table(table: list[list], limit: int = 3) -> str:
    """The first few non-empty rows, for an error the reader can act on."""
    shown: list[str] = []
    for row in table[:MAX_HEADER_SCAN]:
        cells = [("" if c is None else str(c).strip()) for c in row]
        if not any(cells):
            continue
        text = " | ".join(c for c in cells if c)[:200]
        shown.append(text)
        if len(shown) >= limit:
            break
    if not shown:
        return "the file contains no non-empty rows"
    return "; ".join(f'"{s}"' for s in shown)


def _rows_from_table(
    table: list[list], source_name: str, sheet: str = ""
) -> RawRecipe:
    header_index = find_header(table)
    if header_index is None:
        where = f" on sheet {sheet!r}" if sheet else ""
        raise RecipeLoadError(
            "No recognisable header row was found"
            f"{where}. The recipe needs a row naming the columns Platform, "
            "Category, Component and Target Version. The first rows found were: "
            f"{describe_table(table)}."
        )

    header = ["" if c is None else str(c).strip() for c in table[header_index]]
    mapping = _map_header(header)

    rows: list[dict] = []
    for offset, raw_row in enumerate(table[header_index + 1 :], start=1):
        if len(rows) >= MAX_ROWS:
            raise RecipeLoadError(f"The recipe exceeds the {MAX_ROWS}-row limit.")
        record: dict = {"_row": header_index + 1 + offset, "_raw": {}}
        has_value = False
        for col_index, column in enumerate(header):
            value = raw_row[col_index] if col_index < len(raw_row) else None
            text = "" if value is None else str(value).strip()
            if column:
                record["_raw"][column] = text
            if text:
                has_value = True
            target = mapping.get(column)
            if target:
                record[target] = text
        record["_empty"] = not has_value
        rows.append(record)

    return RawRecipe(
        rows=rows,
        columns=[c for c in header if c],
        mapped_columns=mapping,
        source_name=source_name,
        sheet=sheet,
    )


def _load_csv(data: bytes, source_name: str) -> RawRecipe:
    text = data.decode("utf-8-sig", errors="replace")
    sample = text[:8192]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=",;\t|")
    except csv.Error:
        dialect = csv.excel
    reader = csv.reader(io.StringIO(text), dialect)
    table = [row for row in reader]
    if not table:
        raise RecipeLoadError("The recipe file is empty.")
    return _rows_from_table(table, source_name)


def _try_matrix(
    tables: list[tuple[str, list[list]]], source_name: str, continuous: bool = False
):
    """Attempt the vendor version-matrix layout.  Returns None if it isn't one."""
    from .matrix_loader import parse_matrix

    rows, skipped = parse_matrix(tables, continuous=continuous)
    if not rows:
        return None
    sheets = sorted({row["_sheet"] for row in rows})
    return RawRecipe(
        rows=rows,
        columns=["Platform", "Model", "Category", "Component", "Target Version",
                 "Target Build"],
        mapped_columns={
            "Platform": "platform",
            "Model": "model",
            "Category": "category",
            "Component": "component",
            "Target Version": "target_version",
            "Target Build": "target_build",
        },
        source_name=source_name,
        sheet=", ".join(sheets),
        layout="matrix",
        notes=skipped,
    )


def _load_xlsx(path: Path, source_name: str) -> RawRecipe:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:  # pragma: no cover - dependency is declared
        raise RecipeLoadError(
            "XLSX support requires the 'openpyxl' package to be installed."
        ) from exc

    try:
        workbook = load_workbook(path, read_only=True, data_only=True)
    except Exception as exc:
        raise RecipeLoadError(
            "The XLSX file could not be opened. Confirm it is a valid, "
            "uncorrupted Excel workbook."
        ) from exc

    # The recipe is not always on the first sheet, so every sheet is scanned
    # and the one whose header matches best wins.  Named sheets are preferred
    # only as a tie-breaker.
    preferred_names = {"recipe", "recipe data", "firmware", "targets"}
    try:
        tables: list[tuple[str, list[list]]] = []
        for name in workbook.sheetnames:
            worksheet = workbook[name]
            tables.append(
                (name, [list(row) for row in worksheet.iter_rows(values_only=True)])
            )
    finally:
        workbook.close()

    if not tables:
        raise RecipeLoadError("The workbook contains no worksheets.")

    candidates = []
    for name, table in tables:
        index = find_header(table)
        if index is None:
            continue
        cells = ["" if c is None else str(c) for c in table[index]]
        candidates.append((_header_score(cells), norm_text(name) in preferred_names, name, table))

    if not candidates:
        # Not a plain table - try the vendor "version matrix" layout before
        # giving up.
        matrix = _try_matrix(tables, source_name)
        if matrix is not None:
            return matrix
        summary = "; ".join(
            f"sheet {name!r} starts with {describe_table(table, limit=2)}"
            for name, table in tables[:4]
        )
        raise RecipeLoadError(
            "No recognisable header row was found on any worksheet, and the "
            "sheets do not use the vendor version-matrix layout either. The "
            "recipe needs a row naming the columns Platform, Category, "
            f"Component and Target Version. Found: {summary}."
        )

    candidates.sort(key=lambda c: (c[0], c[1]), reverse=True)
    _score, _preferred, name, table = candidates[0]
    return _rows_from_table(table, source_name, sheet=name)


def _load_json(data: bytes, source_name: str) -> RawRecipe:
    try:
        payload = json.loads(data.decode("utf-8-sig", errors="replace"))
    except (ValueError, UnicodeDecodeError) as exc:
        raise RecipeLoadError("The JSON recipe could not be parsed.") from exc

    if isinstance(payload, dict):
        for key in ("recipe", "rows", "entries", "data"):
            if isinstance(payload.get(key), list):
                payload = payload[key]
                break
    if not isinstance(payload, list) or not payload:
        raise RecipeLoadError(
            "The JSON recipe must be a non-empty list of objects, or an object "
            "with a 'recipe' list."
        )
    if not all(isinstance(item, dict) for item in payload):
        raise RecipeLoadError("Every entry in a JSON recipe must be an object.")

    columns: list[str] = []
    for item in payload:
        for key in item:
            if key not in columns:
                columns.append(str(key))
    table = [columns]
    for item in payload:
        table.append([item.get(column, "") for column in columns])
    return _rows_from_table(table, source_name)


_LINE_TOLERANCE = 3.0   # points; words this close vertically are one row
_COLUMN_GAP = 12.0      # points; a horizontal gap this wide starts a column


def _rows_from_words(page) -> list[list]:
    """Rebuild rows and columns from word positions.

    Vendor matrices are often published without ruled borders, so
    ``extract_tables`` returns nothing. Splitting the plain text on runs of
    whitespace does not work either - "Components Recipe 2026.03" is a single
    space apart - so columns are recovered from the horizontal gaps between
    words instead.
    """
    try:
        words = page.extract_words() or []
    except Exception:  # pragma: no cover - defensive
        return []
    if not words:
        return []

    lines: dict[float, list[dict]] = {}
    for word in words:
        top = round(float(word["top"]) / _LINE_TOLERANCE) * _LINE_TOLERANCE
        lines.setdefault(top, []).append(word)

    rows: list[list] = []
    for top in sorted(lines):
        ordered = sorted(lines[top], key=lambda w: float(w["x0"]))
        cells: list[str] = []
        current = [ordered[0]["text"]]
        previous_end = float(ordered[0]["x1"])
        for word in ordered[1:]:
            if float(word["x0"]) - previous_end > _COLUMN_GAP:
                cells.append(" ".join(current))
                current = [word["text"]]
            else:
                current.append(word["text"])
            previous_end = float(word["x1"])
        cells.append(" ".join(current))
        if any(c.strip() for c in cells):
            rows.append(cells)
    return rows


def _load_pdf(path: Path, source_name: str) -> RawRecipe:
    """Read the vendor matrix straight from the published PDF.

    Each page's tables are turned into the same row-and-column shape a
    converted workbook has, so the matrix reader handles both identically.
    """
    try:
        import pdfplumber
    except ImportError as exc:
        raise RecipeLoadError(
            "PDF support needs the 'pdfplumber' package. Install it with "
            "'pip install -r requirements.txt', or convert the PDF to XLSX "
            "and upload that instead."
        ) from exc

    tables: list[tuple[str, list[list]]] = []
    try:
        with pdfplumber.open(path) as pdf:
            for index, page in enumerate(pdf.pages, start=1):
                rows: list[list] = []
                for table in page.extract_tables() or []:
                    for row in table:
                        rows.append(["" if c is None else str(c).strip() for c in row])
                if not rows:
                    rows = _rows_from_words(page)
                if rows:
                    tables.append((f"Page {index}", rows))
    except RecipeLoadError:
        raise
    except Exception as exc:
        raise RecipeLoadError(
            "The PDF could not be read. Confirm it is a text-based PDF rather "
            "than a scan, or convert it to XLSX and upload that."
        ) from exc

    if not tables:
        raise RecipeLoadError("No tables or text were found in the PDF.")

    # The pages are one continuous document: a section that starts on one page
    # and runs onto the next is still that section.
    matrix = _try_matrix(tables, source_name, continuous=True)
    if matrix is not None:
        return matrix

    # Not a version matrix - try each page as a plain table.
    for name, rows in tables:
        if find_header(rows) is not None:
            return _rows_from_table(rows, source_name, sheet=name)

    summary = "; ".join(
        f"{name} starts with {describe_table(rows, limit=2)}" for name, rows in tables[:3]
    )
    raise RecipeLoadError(
        "The PDF was read, but no recognisable recipe layout was found. "
        f"Found: {summary}."
    )


def load_recipe(path: str | Path, source_name: str | None = None) -> RawRecipe:
    """Load a recipe from disk.  Raises :class:`RecipeLoadError` on failure."""
    path = Path(path)
    name = source_name or path.name
    suffix = path.suffix.lower()

    if suffix == ".pdf":
        recipe = _load_pdf(path, name)
        logger.info("Loaded recipe %s from PDF: %d row(s)", name, len(recipe.rows))
        return recipe
    if suffix == ".xls":
        raise RecipeLoadError(
            "Legacy .xls workbooks are not supported. Please re-save the recipe "
            "as .xlsx or .csv and upload it again."
        )
    if suffix == ".xlsx":
        recipe = _load_xlsx(path, name)
    elif suffix == ".csv":
        recipe = _load_csv(path.read_bytes(), name)
    elif suffix == ".json":
        recipe = _load_json(path.read_bytes(), name)
    else:
        raise RecipeLoadError(f"Unsupported recipe file type {suffix!r}.")

    logger.info(
        "Loaded recipe %s: %d row(s), columns mapped: %s",
        name,
        len(recipe.rows),
        sorted(set(recipe.mapped_columns.values())),
    )
    return recipe
