"""Multi-worksheet XLSX report.

Sheets: Summary, Comparison, Installed_Data (raw X-axis), Recipe_Data (raw
Y-axis) and Exceptions (everything that is not compliant), so the whole run is
auditable from one file.
"""
from __future__ import annotations

import io
from datetime import datetime

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from firmware.models import (
    COMPARISON_COLUMNS,
    RAW_COLUMNS,
    RECIPE_COLUMNS,
    STATUS_NEEDS_UPDATE,
    STATUS_NOT_APPLICABLE,
    STATUS_NOT_DETECTED,
    STATUS_NOT_IN_RECIPE,
    STATUS_UNABLE,
    STATUS_UPDATED,
)
from .summary import compliance_rate

_HEADER_FILL = PatternFill("solid", fgColor="1F3864")
_HEADER_FONT = Font(color="FFFFFF", bold=True)
_TITLE_FONT = Font(size=14, bold=True)

_STATUS_FILL = {
    STATUS_UPDATED: PatternFill("solid", fgColor="D5F5E3"),
    STATUS_NEEDS_UPDATE: PatternFill("solid", fgColor="FADBD8"),
    STATUS_NOT_DETECTED: PatternFill("solid", fgColor="FCF3CF"),
    STATUS_NOT_IN_RECIPE: PatternFill("solid", fgColor="EAECEE"),
    STATUS_UNABLE: PatternFill("solid", fgColor="FDEBD0"),
    STATUS_NOT_APPLICABLE: PatternFill("solid", fgColor="EAECEE"),
}

_MAX_WIDTH = 60


def _cell(value):
    if value is None:
        return ""
    if isinstance(value, bool):
        return "Yes" if value else "No"
    if isinstance(value, (int, float, str)):
        return value
    return str(value)


def _write_sheet(worksheet, rows, columns, status_key: str | None = None) -> None:
    worksheet.append([label for _key, label in columns])
    for index in range(1, len(columns) + 1):
        cell = worksheet.cell(row=1, column=index)
        cell.fill = _HEADER_FILL
        cell.font = _HEADER_FONT
        cell.alignment = Alignment(vertical="center")

    widths = [len(label) for _key, label in columns]
    for row in rows:
        data = row.to_dict() if hasattr(row, "to_dict") else dict(row)
        values = [_cell(data.get(key)) for key, _label in columns]
        worksheet.append(values)
        for position, value in enumerate(values):
            widths[position] = max(widths[position], min(len(str(value)), _MAX_WIDTH))
        if status_key:
            status = data.get(status_key)
            fill = _STATUS_FILL.get(status)
            if fill:
                column_index = [k for k, _ in columns].index(status_key) + 1
                worksheet.cell(row=worksheet.max_row, column=column_index).fill = fill

    for position, width in enumerate(widths, start=1):
        worksheet.column_dimensions[get_column_letter(position)].width = min(
            width + 2, _MAX_WIDTH
        )
    worksheet.freeze_panes = "A2"
    if rows:
        worksheet.auto_filter.ref = (
            f"A1:{get_column_letter(len(columns))}{worksheet.max_row}"
        )


def _write_summary(worksheet, result) -> None:
    counts = result.summary.get("counts", {})
    worksheet["A1"] = "OpsRamp Firmware Compliance Report"
    worksheet["A1"].font = _TITLE_FONT
    rows = [
        ("Generated", datetime.now().strftime("%Y-%m-%d %H:%M:%S")),
        ("Tenant", result.tenant_name),
        ("Tenant ID", result.tenant_id),
        ("Category", result.summary.get("category", result.category)),
        ("Recipe file", result.recipe_name),
        ("Recipe entries", result.summary.get("recipe_entries", 0)),
        ("", ""),
        ("Assets discovered in tenant", result.assets_total),
        ("Assets in scope", result.summary.get("assets_in_scope", 0)),
        ("Components checked", result.summary.get("components_checked", 0)),
        ("", ""),
        (STATUS_UPDATED, counts.get(STATUS_UPDATED, 0)),
        (STATUS_NEEDS_UPDATE, counts.get(STATUS_NEEDS_UPDATE, 0)),
        (STATUS_NOT_DETECTED, counts.get(STATUS_NOT_DETECTED, 0)),
        (STATUS_NOT_IN_RECIPE, counts.get(STATUS_NOT_IN_RECIPE, 0)),
        (STATUS_UNABLE, counts.get(STATUS_UNABLE, 0)),
        (STATUS_NOT_APPLICABLE, counts.get(STATUS_NOT_APPLICABLE, 0)),
        ("", ""),
        ("Compliance rate (updated / comparable)", f"{compliance_rate(result)}%"),
    ]
    for offset, (label, value) in enumerate(rows, start=3):
        worksheet.cell(row=offset, column=1, value=label).font = Font(bold=bool(label))
        worksheet.cell(row=offset, column=2, value=_cell(value))
        fill = _STATUS_FILL.get(label)
        if fill:
            worksheet.cell(row=offset, column=1).fill = fill
    worksheet.column_dimensions["A"].width = 40
    worksheet.column_dimensions["B"].width = 42


def build_workbook(result) -> bytes:
    workbook = Workbook()

    summary_sheet = workbook.active
    summary_sheet.title = "Summary"
    _write_summary(summary_sheet, result)

    _write_sheet(
        workbook.create_sheet("Comparison"), result.rows, COMPARISON_COLUMNS, "result"
    )
    _write_sheet(
        workbook.create_sheet("Installed_Data"), result.raw_rows, RAW_COLUMNS
    )
    _write_sheet(
        workbook.create_sheet("Recipe_Data"), result.recipe_entries, RECIPE_COLUMNS
    )

    exceptions = [
        row
        for row in result.rows
        if row.result
        in (
            STATUS_NEEDS_UPDATE,
            STATUS_NOT_DETECTED,
            STATUS_NOT_IN_RECIPE,
            STATUS_UNABLE,
        )
    ]
    _write_sheet(
        workbook.create_sheet("Exceptions"), exceptions, COMPARISON_COLUMNS, "result"
    )

    buffer = io.BytesIO()
    workbook.save(buffer)
    workbook.close()
    return buffer.getvalue()
