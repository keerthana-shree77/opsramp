"""Convert the CSV recipe template into an XLSX workbook.

Usage:
    python tools/make_recipe_xlsx.py [source.csv] [target.xlsx]
"""
from __future__ import annotations

import csv
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parent.parent
DEFAULT_SOURCE = ROOT / "static" / "recipe_template.csv"
DEFAULT_TARGET = ROOT / "static" / "recipe_template.xlsx"


def main() -> int:
    source = Path(sys.argv[1]) if len(sys.argv) > 1 else DEFAULT_SOURCE
    target = Path(sys.argv[2]) if len(sys.argv) > 2 else DEFAULT_TARGET

    with source.open(newline="", encoding="utf-8-sig") as handle:
        rows = list(csv.reader(handle))
    if not rows:
        print(f"{source} is empty")
        return 1

    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Recipe"
    for row in rows:
        worksheet.append(row)

    header_fill = PatternFill("solid", fgColor="1F3864")
    for index in range(1, len(rows[0]) + 1):
        cell = worksheet.cell(row=1, column=index)
        cell.fill = header_fill
        cell.font = Font(color="FFFFFF", bold=True)
        width = max(len(str(r[index - 1])) for r in rows if len(r) >= index)
        worksheet.column_dimensions[get_column_letter(index)].width = min(width + 3, 42)
    worksheet.freeze_panes = "A2"

    workbook.save(target)
    print(f"Wrote {target}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
