"""Print the raw layout of a recipe workbook so its structure can be read.

Usage:
    python tools/inspect_recipe.py <file.xlsx|file.csv> [output.txt]

Prints, for every worksheet, the first rows exactly as they appear - including
empty rows and merged-cell gaps - so a layout that the loader did not
recognise can be diagnosed without guessing.
"""
from __future__ import annotations

import csv
import sys
from pathlib import Path

MAX_ROWS = 400
MAX_COLS = 14
MAX_CELL = 60


def cell(value) -> str:
    if value is None:
        return ""
    text = str(value).strip().replace("\n", " ")
    if len(text) > MAX_CELL:
        text = text[: MAX_CELL - 3] + "..."
    return text


def dump_table(name: str, rows: list[list], out: list[str]) -> None:
    out.append("")
    out.append("=" * 72)
    out.append(f"SHEET: {name}   ({len(rows)} rows total)")
    out.append("=" * 72)
    shown = 0
    for index, row in enumerate(rows, start=1):
        cells = [cell(c) for c in list(row)[:MAX_COLS]]
        while cells and not cells[-1]:
            cells.pop()
        if not cells:
            out.append(f"  row {index:>3}: (empty)")
        else:
            out.append(f"  row {index:>3}: " + " | ".join(cells))
        shown += 1
        if shown >= MAX_ROWS:
            remaining = len(rows) - shown
            if remaining > 0:
                out.append(f"  ... {remaining} more rows not shown")
            break


def main() -> int:
    if len(sys.argv) < 2:
        print("Usage: python tools/inspect_recipe.py <file.xlsx|file.csv> [output.txt]")
        return 1

    path = Path(sys.argv[1])
    if not path.exists():
        print(f"File not found: {path}")
        return 1

    out: list[str] = [
        "RECIPE FILE LAYOUT REPORT",
        f"File: {path.name}",
        f"Size: {path.stat().st_size:,} bytes",
    ]

    suffix = path.suffix.lower()
    if suffix in (".xlsx", ".xlsm"):
        from openpyxl import load_workbook

        workbook = load_workbook(path, read_only=True, data_only=True)
        try:
            out.append(f"Worksheets: {', '.join(workbook.sheetnames)}")
            for name in workbook.sheetnames:
                rows = [list(r) for r in workbook[name].iter_rows(values_only=True)]
                dump_table(name, rows, out)
        finally:
            workbook.close()
    elif suffix in (".csv", ".txt"):
        with path.open(newline="", encoding="utf-8-sig", errors="replace") as handle:
            rows = list(csv.reader(handle))
        dump_table(path.name, rows, out)
    else:
        print(f"Unsupported file type: {suffix}. Use .xlsx or .csv")
        return 1

    report = "\n".join(out)
    if len(sys.argv) > 2:
        Path(sys.argv[2]).write_text(report, encoding="utf-8")
        print(f"Wrote {sys.argv[2]}")
    else:
        print(report)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
